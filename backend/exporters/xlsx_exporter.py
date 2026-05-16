from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from schemas.candidate import ComparisonReport, ComplianceReport


def _header_cell(cell, text: str) -> None:
    cell.value = text
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F2937")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_comparison_xlsx(
    comparison: ComparisonReport,
    compliance: ComplianceReport,
    path: Path,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"

    must_have_skills: list[str] = []
    seen: set[str] = set()
    for row in comparison.rows:
        for skill in row.must_have_ratings.keys():
            if skill not in seen:
                seen.add(skill)
                must_have_skills.append(skill)

    headers = ["Candidate"] + must_have_skills + ["Red Flag", "Recommended Focus"]
    for col, header in enumerate(headers, start=1):
        _header_cell(ws.cell(row=1, column=col), header)

    for r, row in enumerate(comparison.rows, start=2):
        ws.cell(row=r, column=1, value=row.candidate_name)
        for c, skill in enumerate(must_have_skills, start=2):
            ws.cell(row=r, column=c, value=row.must_have_ratings.get(skill, "—"))
        ws.cell(row=r, column=len(must_have_skills) + 2, value=row.red_flag_status)
        ws.cell(
            row=r,
            column=len(must_have_skills) + 3,
            value="; ".join(row.recommended_focus_areas),
        )

    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 22

    ws.cell(row=len(comparison.rows) + 3, column=1, value="AI-assisted analysis — recruiter review required.").font = Font(italic=True)

    cs = wb.create_sheet("Compliance Gaps")
    _header_cell(cs.cell(row=1, column=1), "Candidate")
    _header_cell(cs.cell(row=1, column=2), "Missing Required Checks")
    for r, gap in enumerate(compliance.gaps, start=2):
        cs.cell(row=r, column=1, value=gap.candidate_name)
        cs.cell(row=r, column=2, value=", ".join(gap.missing_checks) if gap.missing_checks else "None")
    cs.column_dimensions["A"].width = 28
    cs.column_dimensions["B"].width = 60

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
