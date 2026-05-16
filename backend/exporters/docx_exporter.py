from pathlib import Path

from docx import Document

from schemas.feedback import FeedbackTemplate


def write_feedback_docx(template: FeedbackTemplate, path: Path) -> None:
    doc = Document()
    doc.add_heading(f"Panelist Feedback — {template.candidate_name}", level=1)
    doc.add_paragraph(f"Role: {template.role_title}")

    disclaimer = doc.add_paragraph()
    run = disclaimer.add_run("AI-assisted analysis — recruiter review required.")
    run.italic = True

    doc.add_heading("Rating Scale", level=2)
    for label in template.rating_scale:
        doc.add_paragraph(label, style="List Bullet")

    doc.add_heading("Competency Ratings", level=2)
    for comp in template.competencies:
        doc.add_heading(comp.competency, level=3)
        if comp.definition:
            doc.add_paragraph(f"Definition: {comp.definition}")
        doc.add_paragraph("Rating (circle one): " + " / ".join(template.rating_scale))
        doc.add_paragraph("Observations:")
        doc.add_paragraph("________________________________________")
        doc.add_paragraph("________________________________________")

    doc.add_heading("Red Flag Indicators", level=2)
    if template.red_flag_indicators:
        for flag in template.red_flag_indicators:
            doc.add_paragraph(f"[ ] {flag}")
    else:
        doc.add_paragraph("None defined.")

    doc.add_heading("Compliance Checklist", level=2)
    if template.compliance_checks:
        for check in template.compliance_checks:
            doc.add_paragraph(f"[ ] {check}")
    else:
        doc.add_paragraph("None defined.")

    doc.add_heading("Overall Recommendation", level=2)
    doc.add_paragraph("[ ] Proceed   [ ] Hold   [ ] Decline")
    doc.add_paragraph("Justification (required):")
    doc.add_paragraph("________________________________________")
    doc.add_paragraph("________________________________________")

    doc.save(str(path))
